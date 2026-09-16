"""Convert uploaded documents (PDF/Word/Excel/HTML) to Markdown for KB display."""
import os
import re


def _table_to_md(rows):
    rows = [[("" if c is None else str(c)).replace("\n", " ").strip() for c in r] for r in rows]
    rows = [r for r in rows if any(c for c in r)]
    if not rows:
        return ""
    width = max(len(r) for r in rows)
    rows = [r + [""] * (width - len(r)) for r in rows]
    out = ["| " + " | ".join(rows[0]) + " |", "| " + " | ".join(["---"] * width) + " |"]
    for r in rows[1:]:
        out.append("| " + " | ".join(r) + " |")
    return "\n".join(out)


def pdf_to_md(path: str) -> str:
    import pdfplumber
    parts = []
    with pdfplumber.open(path) as pdf:
        for i, page in enumerate(pdf.pages):
            txt = page.extract_text() or ""
            parts.append(txt)
            for tb in (page.extract_tables() or []):
                md = _table_to_md(tb)
                if md:
                    parts.append("\n" + md + "\n")
            if i > 200:
                break
    return "\n\n".join(p for p in parts if p.strip())


def docx_to_md(path: str) -> str:
    import docx
    doc = docx.Document(path)
    out = []
    for block in doc.paragraphs:
        style = (block.style.name or "").lower()
        text = block.text.strip()
        if not text:
            continue
        if style.startswith("heading"):
            try:
                lvl = int(re.sub(r"\D", "", style) or "1")
            except ValueError:
                lvl = 1
            out.append("#" * min(lvl, 6) + " " + text)
        elif "list" in style or block.text.lstrip().startswith(("•", "·")):
            out.append("- " + text.lstrip("•· "))
        else:
            out.append(text)
    for tbl in doc.tables:
        rows = [[c.text for c in row.cells] for row in tbl.rows]
        md = _table_to_md(rows)
        if md:
            out.append("\n" + md + "\n")
    return "\n\n".join(out)


def xlsx_to_md(path: str) -> str:
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    out = []
    for ws in wb.worksheets:
        out.append("## %s" % ws.title)
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        md = _table_to_md(rows)
        if md:
            out.append(md)
    return "\n\n".join(out)


def html_to_md(path: str) -> str:
    with open(path, "r", encoding="utf-8", errors="ignore") as fh:
        raw = fh.read()
    raw = re.sub(r"(?is)<(script|style).*?</\1>", " ", raw)
    raw = re.sub(r"(?is)<br\s*/?>", "\n", raw)
    raw = re.sub(r"(?is)</(p|div|tr|li|h[1-6])>", "\n", raw)
    txt = re.sub(r"(?s)<[^>]+>", "", raw)
    txt = re.sub(r"\n{3,}", "\n\n", txt)
    return txt.strip()


def bytes_to_md(name: str, data: bytes) -> str:
    ext = os.path.splitext(name)[1].lower()
    tmp = os.path.join(os.path.dirname(__file__), "data", "_conv_" + str(abs(hash(name + str(len(data))))) + ext)
    os.makedirs(os.path.dirname(tmp), exist_ok=True)
    try:
        with open(tmp, "wb") as fh:
            fh.write(data)
        return convert_file(name, tmp)
    finally:
        try:
            os.remove(tmp)
        except OSError:
            pass


def convert_file(name: str, path: str) -> str:
    ext = os.path.splitext(name)[1].lower()
    try:
        if ext == ".pdf":
            return pdf_to_md(path)
        if ext in (".docx", ".doc"):
            try:
                return docx_to_md(path)
            except Exception:
                with open(path, "r", encoding="utf-8", errors="ignore") as fh:
                    return fh.read()
        if ext in (".xlsx", ".xls"):
            return xlsx_to_md(path)
        if ext in (".md", ".markdown", ".txt"):
            with open(path, "r", encoding="utf-8", errors="ignore") as fh:
                return fh.read()
        if ext in (".html", ".htm"):
            return html_to_md(path)
    except Exception:
        return ""
    return ""
